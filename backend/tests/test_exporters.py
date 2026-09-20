"""Tests for backend/exporters/excel_exporter.py and json_exporter.py"""

import io
import json

import openpyxl
import pytest

from backend.models.job import Job, ScoredJob
from backend.exporters.excel_exporter import export_excel, _experience
from backend.exporters.json_exporter import export_json


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _job(**kwargs) -> Job:
    defaults = dict(
        id="1", title="Python Developer", company="Acme", location="Bangkok",
        salary_min=None, salary_max=None, currency="THB",
        salary_display=None, salary_is_numeric=False,
        work_arrangement="Hybrid",
        experience_min_years=None, experience_max_years=None,
        categories=[], skills_mentioned=[],
        abstract=None, content=None,
        url="https://example.com/job/1",
        listed_at="2025-01-01T00:00:00+00:00",
        source="jobsdb", raw=None,
    )
    defaults.update(kwargs)
    return Job(**defaults)


def _scored(rank=1, score=0.85, **job_kwargs) -> ScoredJob:
    return ScoredJob(job=_job(**job_kwargs), score=score, rank=rank)


def _parse_excel(data: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(data))


# ---------------------------------------------------------------------------
# _experience helper
# ---------------------------------------------------------------------------

class TestExperienceHelper:
    def test_both_bounds(self):
        assert _experience(3.0, 7.0) == "3-7 yrs"

    def test_min_only(self):
        assert _experience(1.0, None) == "1+ yrs"

    def test_both_none(self):
        assert _experience(None, None) == "Not specified"

    def test_zero_min(self):
        assert _experience(0.0, 2.0) == "0-2 yrs"

    def test_strips_trailing_dot_zero(self):
        # :g format — 3.0 should render as "3" not "3.0"
        assert _experience(3.0, 5.0) == "3-5 yrs"
        assert _experience(5.0, None) == "5+ yrs"


# ---------------------------------------------------------------------------
# export_excel — return type and valid xlsx
# ---------------------------------------------------------------------------

class TestExportExcelOutput:
    def test_returns_bytes(self):
        result = export_excel([_scored()])
        assert isinstance(result, bytes)

    def test_returns_nonempty(self):
        assert len(export_excel([_scored()])) > 0

    def test_valid_xlsx_parseable(self):
        # openpyxl.load_workbook raises if bytes are corrupt
        wb = _parse_excel(export_excel([_scored()]))
        assert wb.active is not None

    def test_empty_list_still_valid(self):
        wb = _parse_excel(export_excel([]))
        ws = wb.active
        assert ws.max_row == 1  # header only


# ---------------------------------------------------------------------------
# export_excel — header row
# ---------------------------------------------------------------------------

class TestExportExcelHeader:
    EXPECTED = [
        "Rank", "Score (%)", "Title", "Company", "Location",
        "Salary", "Work Arrangement", "Experience Level", "Category", "URL", "Listed At",
    ]

    def setup_method(self):
        self.ws = _parse_excel(export_excel([_scored()])).active

    def test_header_values(self):
        actual = [self.ws.cell(1, c).value for c in range(1, 12)]
        assert actual == self.EXPECTED

    def test_header_bold(self):
        for col in range(1, 12):
            assert self.ws.cell(1, col).font.bold

    def test_header_fill_colour(self):
        for col in range(1, 12):
            fill = self.ws.cell(1, col).fill
            assert fill.fgColor.rgb.upper().endswith("BDD7EE")


# ---------------------------------------------------------------------------
# export_excel — data row values
# ---------------------------------------------------------------------------

class TestExportExcelDataRows:
    def _row(self, **kwargs):
        ws = _parse_excel(export_excel([_scored(**kwargs)])).active
        return [ws.cell(2, c).value for c in range(1, 12)]

    def test_rank_and_score(self):
        row = self._row(rank=3, score=0.72)
        assert row[0] == 3       # Rank
        assert row[1] == pytest.approx(0.72)  # Score stored as float

    def test_title(self):
        row = self._row(title="Senior Dev")
        assert row[2] == "Senior Dev"

    def test_company_location(self):
        row = self._row(company="TechCo", location="Chiang Mai")
        assert row[3] == "TechCo"
        assert row[4] == "Chiang Mai"

    def test_salary_from_display(self):
        row = self._row(salary_display="80,000-100,000")
        assert row[5] == "80,000-100,000"

    def test_salary_none_shows_other(self):
        row = self._row(salary_display=None)
        assert row[5] == "Other"

    def test_salary_empty_string_shows_other(self):
        row = self._row(salary_display="")
        assert row[5] == "Other"

    def test_work_arrangement(self):
        row = self._row(work_arrangement="Remote")
        assert row[6] == "Remote"

    def test_experience_both_bounds(self):
        row = self._row(experience_min_years=2.0, experience_max_years=5.0)
        assert row[7] == "2-5 yrs"

    def test_experience_min_only(self):
        row = self._row(experience_min_years=3.0, experience_max_years=None)
        assert row[7] == "3+ yrs"

    def test_experience_not_specified(self):
        row = self._row(experience_min_years=None, experience_max_years=None)
        assert row[7] == "Not specified"

    def test_categories_joined(self):
        row = self._row(categories=["IT", "Backend"])
        assert row[8] == "IT, Backend"

    def test_categories_empty(self):
        row = self._row(categories=[])
        # openpyxl stores empty string as None on round-trip
        assert not row[8]

    def test_url(self):
        row = self._row()
        assert row[9] == "https://example.com/job/1"

    def test_listed_at(self):
        row = self._row(listed_at="2025-06-01T00:00:00+00:00")
        assert row[10] == "2025-06-01T00:00:00+00:00"

    def test_none_fields_not_blank_crash(self):
        # None fields write empty string; openpyxl round-trips empty string as None
        row = self._row(company=None, location=None, work_arrangement=None)
        assert not row[3]
        assert not row[4]
        assert not row[6]


# ---------------------------------------------------------------------------
# export_excel — score number format
# ---------------------------------------------------------------------------

class TestExportExcelScoreFormat:
    def test_score_number_format_is_percentage(self):
        ws = _parse_excel(export_excel([_scored(score=0.85)])).active
        fmt = ws.cell(2, 2).number_format
        assert "%" in fmt

    def test_score_one_decimal(self):
        ws = _parse_excel(export_excel([_scored(score=0.85)])).active
        fmt = ws.cell(2, 2).number_format
        assert "0.0%" in fmt


# ---------------------------------------------------------------------------
# export_excel — title hyperlink
# ---------------------------------------------------------------------------

class TestExportExcelHyperlink:
    def test_title_cell_has_hyperlink(self):
        ws = _parse_excel(export_excel([_scored()])).active
        cell = ws.cell(2, 3)
        assert cell.hyperlink is not None

    def test_hyperlink_target_matches_url(self):
        ws = _parse_excel(export_excel([_scored(url="https://jobs.example.com/42")])).active
        cell = ws.cell(2, 3)
        assert cell.hyperlink.target == "https://jobs.example.com/42"


# ---------------------------------------------------------------------------
# export_excel — column widths
# ---------------------------------------------------------------------------

class TestExportExcelColumnWidth:
    def test_columns_have_positive_width(self):
        from openpyxl.utils import get_column_letter
        ws = _parse_excel(export_excel([_scored()])).active
        for col in range(1, 12):
            letter = get_column_letter(col)
            dim = ws.column_dimensions[letter]
            assert dim.width > 0

    def test_width_capped_at_62(self):
        # 60 char cap + 2 padding = 62 max
        from openpyxl.utils import get_column_letter
        long_title = "A" * 200
        ws = _parse_excel(export_excel([_scored(title=long_title)])).active
        for col in range(1, 12):
            letter = get_column_letter(col)
            assert ws.column_dimensions[letter].width <= 62


# ---------------------------------------------------------------------------
# export_excel — row count
# ---------------------------------------------------------------------------

class TestExportExcelRowCount:
    def test_n_jobs_produce_n_plus_one_rows(self):
        jobs = [_scored(rank=i + 1, score=0.5, id=str(i)) for i in range(5)]
        ws = _parse_excel(export_excel(jobs)).active
        assert ws.max_row == 6  # 1 header + 5 data


# ---------------------------------------------------------------------------
# export_json
# ---------------------------------------------------------------------------

class TestExportJson:
    def test_returns_bytes(self):
        assert isinstance(export_json([], "abc"), bytes)

    def test_valid_json(self):
        data = json.loads(export_json([_scored()], "sid-1"))
        assert data["search_id"] == "sid-1"
        assert data["total"] == 1
        assert len(data["jobs"]) == 1

    def test_empty_list(self):
        data = json.loads(export_json([], "sid-x"))
        assert data["total"] == 0
        assert data["jobs"] == []

    def test_exported_at_present(self):
        data = json.loads(export_json([], "sid-x"))
        assert "exported_at" in data

    def test_utf8_encoding(self):
        result = export_json([_scored(title="นักพัฒนา")], "s")
        data = json.loads(result.decode("utf-8"))
        assert data["jobs"][0]["job"]["title"] == "นักพัฒนา"
