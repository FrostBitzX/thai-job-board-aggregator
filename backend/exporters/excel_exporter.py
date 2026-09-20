import io
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Allow running as __main__ directly (adds project root to path before package import)
if __name__ == "__main__":
    import pathlib
    sys.path.insert(0, str(pathlib.Path(__file__).parents[2]))

from backend.models.job import ScoredJob

_HEADERS = [
    "Rank", "Score (%)", "Title", "Company", "Location",
    "Salary", "Work Arrangement", "Experience Level", "Category", "URL", "Listed At",
]
_HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
_HEADER_FONT = Font(bold=True)
_MAX_COL_WIDTH = 60


def _experience(min_yrs: float | None, max_yrs: float | None) -> str:
    if min_yrs is None and max_yrs is None:
        return "Not specified"
    if max_yrs is None:
        return f"{min_yrs:g}+ yrs"
    return f"{min_yrs:g}-{max_yrs:g} yrs"


def export_excel(jobs: list[ScoredJob]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active

    # Header row
    for col, header in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL

    # Track max content width per column (start with header widths)
    col_widths = [len(h) for h in _HEADERS]

    for row_idx, sj in enumerate(jobs, 2):
        j = sj.job

        salary = j.salary_display if j.salary_display else "Other"
        experience = _experience(j.experience_min_years, j.experience_max_years)
        category = ", ".join(j.categories)

        row_values = [
            sj.rank,
            sj.score,           # stored as float; formatted as 0.0% via number_format
            j.title,            # gets hyperlink below
            j.company or "",
            j.location or "",
            salary,
            j.work_arrangement or "",
            experience,
            category,
            j.url,
            j.listed_at or "",
        ]

        for col_idx, value in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            col_widths[col_idx - 1] = min(
                _MAX_COL_WIDTH,
                max(col_widths[col_idx - 1], len(str(value)))
            )

        # Score: percentage format
        ws.cell(row=row_idx, column=2).number_format = "0.0%"

        # Title: hyperlink
        title_cell = ws.cell(row=row_idx, column=3)
        title_cell.hyperlink = j.url
        title_cell.style = "Hyperlink"

    # Apply column widths (+2 padding)
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width + 2

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    from backend.models.job import Job, ScoredJob

    def _make_job(rank, score, title, salary_display, exp_min, exp_max, categories):
        j = Job(
            id=str(rank),
            title=title,
            company="Acme Co",
            location="Bangkok",
            salary_min=None,
            salary_max=None,
            currency="THB",
            salary_display=salary_display,
            salary_is_numeric=False,
            work_arrangement="Hybrid",
            experience_min_years=exp_min,
            experience_max_years=exp_max,
            categories=categories,
            skills_mentioned=[],
            abstract=None,
            content=None,
            url=f"https://example.com/job/{rank}",
            listed_at="2025-01-01T00:00:00+00:00",
            source="jobsdb",
            raw=None,
        )
        return ScoredJob(job=j, score=score, rank=rank)

    fake_jobs = [
        _make_job(1, 0.85, "Senior Python Dev", "80,000-100,000", 3.0, 7.0, ["IT", "Backend"]),
        _make_job(2, 0.72, "Data Analyst", None, 1.0, None, []),
        _make_job(3, 0.60, "Frontend Engineer", "ตามตกลง", None, None, ["IT"]),
    ]

    result = export_excel(fake_jobs)
    assert len(result) > 0, "export_excel returned empty bytes"
    print("self-check passed")
